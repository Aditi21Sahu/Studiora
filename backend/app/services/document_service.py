import os
import re
import logging
from typing import Tuple

logger = logging.getLogger(__name__)

class DocumentService:
    """Service to safely parse and extract readable educational text from PDF and DOCX files."""

    @staticmethod
    def clean_text(text: str) -> str:
        """Normalize whitespaces, strip null bytes and format carriage returns."""
        if not text:
            return ""
        # Remove null characters
        cleaned = text.replace("\x00", "")
        # Normalize carriage returns
        cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
        # Collapse excessive newlines
        cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
        # Collapse multiple horizontal spaces
        cleaned = re.sub(r'[ \t]+', ' ', cleaned)
        return cleaned.strip()

    @classmethod
    def extract_text_from_pdf(cls, file_path: str) -> Tuple[str, str]:
        """
        Extract text from a PDF file using PyMuPDF (fitz) with fallback to pypdf.
        Returns: (extracted_text, size_or_page_summary)
        Raises: ValueError if the PDF is unparseable or contains no extractable text.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PDF file not found at: {file_path}")

        full_text = ""
        num_pages = 0

        # 1. Try PyMuPDF (fitz) - high accuracy
        try:
            import fitz
            doc = fitz.open(file_path)
            num_pages = len(doc)
            extracted_pages = []
            for page in doc:
                text = page.get_text("text")
                if text and text.strip():
                    extracted_pages.append(text.strip())
            doc.close()
            full_text = "\n\n".join(extracted_pages)
        except Exception as e:
            logger.warning("PyMuPDF extraction failed, trying pypdf: %s", str(e))
            # 2. Fallback to pypdf
            try:
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                num_pages = len(reader.pages)
                extracted_pages = [p.extract_text().strip() for p in reader.pages if p.extract_text()]
                full_text = "\n\n".join(extracted_pages)
            except Exception as pe:
                logger.error("Error reading PDF %s: %s", file_path, str(pe))
                raise ValueError(f"Unable to parse PDF document: {str(pe)}")

        cleaned = cls.clean_text(full_text)
        if not cleaned or len(cleaned.strip()) < 20:
            raise ValueError("This PDF does not contain selectable text. OCR processing is required.")

        summary = f"{num_pages} {'page' if num_pages == 1 else 'pages'}"
        return cleaned, summary

    @classmethod
    def extract_text_from_docx(cls, file_path: str) -> Tuple[str, str]:
        """
        Extract text from a DOCX file using python-docx.
        Returns: (extracted_text, size_or_page_summary)
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found at: {file_path}")

        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        paragraphs.append(" | ".join(row_text))

            full_text = "\n\n".join(paragraphs)
            cleaned = cls.clean_text(full_text)
            if not cleaned or len(cleaned.strip()) < 20:
                raise ValueError("Could not extract readable text from this Word document. Please ensure the document is not empty.")

            summary = f"{len(paragraphs)} paragraphs"
            return cleaned, summary
        except ValueError:
            raise
        except Exception as e:
            logger.error("Error reading DOCX %s: %s", file_path, str(e))
            raise ValueError(f"Unable to parse Word document: {str(e)}")

document_service = DocumentService()
