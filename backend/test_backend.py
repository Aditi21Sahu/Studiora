import sys
import os
import json

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def run_all_tests():
    print("=== STARTING STUDIORA BACKEND TESTS ===")

    # 1. Health check
    res = client.get("/api/health")
    assert res.status_code == 200, f"Health check failed: {res.text}"
    print("[OK] Health check passed")

    # 2. Register user
    test_user = {
        "full_name": "Aditi Sahu",
        "email": "aditi.test@studiora.com",
        "password": "Password123!",
        "confirm_password": "Password123!"
    }
    res = client.post("/api/auth/register", json=test_user)
    if res.status_code == 400 and "already exists" in res.text:
        # Login instead
        res = client.post("/api/auth/login", json={"email": test_user["email"], "password": test_user["password"]})
    assert res.status_code in [200, 201], f"Auth register/login failed: {res.text}"
    token_data = res.json()
    token = token_data["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print(f"[OK] Auth passed for user {token_data['full_name']}")

    # 3. Check /api/auth/me
    res = client.get("/api/auth/me", headers=headers)
    assert res.status_code == 200
    user_me = res.json()
    assert user_me["email"] == test_user["email"]
    print("[OK] /api/auth/me verified")

    # 4. Check initial dashboard stats
    res = client.get("/api/reports/dashboard-stats", headers=headers)
    assert res.status_code == 200
    stats = res.json()
    print(f"[OK] Initial dashboard stats: materials={stats['study_materials_count']}, notes={stats['notes_count']}, quizzes={stats['quizzes_count']}, avg={stats['average_score']}%")

    # 5. Create a study material (DOCX/PDF simulation or direct material)
    sample_text = (
        "Machine learning is a subfield of artificial intelligence focused on building applications "
        "that learn from data and improve their accuracy over time without being programmed to do so. "
        "In supervised learning, algorithms are trained on labeled data where input-output pairs are provided. "
        "Common supervised algorithms include linear regression, logistic regression, and decision trees. "
        "Unsupervised learning deals with unlabeled data to discover hidden patterns or groupings, such as k-means clustering. "
        "Reinforcement learning trains models using rewards and penalties to achieve optimal goals."
    )
    
    # Test uploading a valid DOCX file via multipart form
    import io
    import docx
    doc = docx.Document()
    doc.add_heading("Machine Learning Fundamentals", 0)
    doc.add_paragraph(sample_text)
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    files = {"file": ("machine_learning_intro.docx", doc_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/materials/upload", files=files, data={"custom_title": "Machine Learning Fundamentals"}, headers=headers)
    assert res.status_code == 201, f"Upload failed: {res.text}"
    material = res.json()
    mat_id = material["id"]
    print(f"[OK] Material uploaded: id={mat_id}, title={material['title']}")

    # 6. Generate notes
    res = client.post(f"/api/notes/generate/{mat_id}", headers=headers)
    assert res.status_code == 201, f"Notes generation failed: {res.text}"
    note = res.json()
    note_id = note["id"]
    assert "title" in note
    assert len(note["sections"]) > 0
    print(f"[OK] Notes generated: id={note_id}, sections={len(note['sections'])}, terms={len(note['important_terms'])}")

    # 7. Test PDF download
    res = client.get(f"/api/notes/{note_id}/download/pdf", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert len(res.content) > 500
    print(f"[OK] PDF download verified ({len(res.content)} bytes)")

    # 8. Test TXT download
    res = client.get(f"/api/notes/{note_id}/download/txt", headers=headers)
    assert res.status_code == 200
    assert "STUDIORA" in res.text
    print("[OK] TXT download verified")

    # 9. Generate 5-question quiz
    res = client.post("/api/quizzes/generate", json={"note_id": note_id, "question_count": 5}, headers=headers)
    assert res.status_code == 201, f"Quiz generation failed: {res.text}"
    quiz = res.json()
    quiz_id = quiz["id"]
    assert len(quiz["questions"]) == 5
    print(f"[OK] Quiz generated: id={quiz_id}, questions={len(quiz['questions'])}")

    # 10. Test checking a single answer interactively
    first_q = quiz["questions"][0]
    res = client.post("/api/quizzes/check-answer", json={
        "quiz_id": quiz_id,
        "question_id": first_q["id"],
        "selected_option": "A"
    }, headers=headers)
    assert res.status_code == 200
    check_data = res.json()
    assert "is_correct" in check_data
    assert "explanation" in check_data
    print(f"[OK] Interactive check-answer response: is_correct={check_data['is_correct']}")

    # 11. Submit full quiz attempt
    answers_payload = [{"question_id": q["id"], "selected_option": "A"} for q in quiz["questions"]]
    res = client.post(f"/api/quizzes/{quiz_id}/submit", json={
        "time_taken_seconds": 45,
        "answers": answers_payload
    }, headers=headers)
    assert res.status_code == 200, f"Quiz submission failed: {res.text}"
    attempt = res.json()
    print(f"[OK] Quiz attempt submitted: score={attempt['score']}/{attempt['total_questions']} ({attempt['percentage']}%)")

    # 12. Check updated dashboard stats
    res = client.get("/api/reports/dashboard-stats", headers=headers)
    assert res.status_code == 200
    updated_stats = res.json()
    assert updated_stats["study_materials_count"] >= 1
    assert updated_stats["notes_count"] >= 1
    assert updated_stats["quizzes_count"] >= 1
    print(f"[OK] Real updated dashboard stats: materials={updated_stats['study_materials_count']}, notes={updated_stats['notes_count']}, quizzes={updated_stats['quizzes_count']}, avg={updated_stats['average_score']}%")

    # 13. Check growth chart
    res = client.get("/api/reports/growth-chart", headers=headers)
    assert res.status_code == 200
    growth = res.json()
    assert len(growth) >= 1
    print(f"[OK] Growth chart verified: {len(growth)} attempts recorded")

    # 14. Check recent activity
    res = client.get("/api/reports/recent-activity", headers=headers)
    assert res.status_code == 200
    activities = res.json()
    assert len(activities) >= 3
    print(f"[OK] Recent activity verified: {len(activities)} activities logged")

    # 15. Verify all frontend static pages & assets
    static_endpoints = [
        "/",
        "/index.html",
        "/login.html",
        "/register.html",
        "/forgot-password.html",
        "/dashboard.html",
        "/study-material.html",
        "/notes.html",
        "/quizzes.html",
        "/reports.html",
        "/css/style.css",
        "/js/api.js",
        "/js/layout.js",
        "/js/dashboard.js",
        "/js/material.js",
        "/js/notes.js",
        "/js/quiz.js",
        "/js/reports.js",
        "/assets/studiora-logo.png",
        "/assets/studiora-student.png",
        "/assets/studiora-dashboard-illustration.png"
    ]
    for endpoint in static_endpoints:
        res = client.get(endpoint)
        assert res.status_code == 200, f"Static route {endpoint} failed with {res.status_code}"
    print(f"[OK] All {len(static_endpoints)} frontend pages and static assets verified (200 OK)")

    print("\n*** ALL BACKEND AND FRONTEND INTEGRATION TESTS PASSED SUCCESSFULLY! ***")

if __name__ == "__main__":
    run_all_tests()

